"""Refine the existing report without changing its SQL, data, or screenshots."""
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / 'CTE_lab_1_completed.docx'
OUTPUT = ROOT / 'output/docs/CTE_lab_1_formatted.docx'
doc = Document(SOURCE)
for style in doc.styles:
    if style.name.lower() in ('heading 1', 'heading 2'):
        style.name = style.name.lower()
original_sql = [p.text for p in doc.paragraphs if p.style.name == 'SQL Code']
original_tables = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]


def font(obj, name='DejaVu Sans', size=11, bold=False, color='000000'):
    obj.font.name = name
    obj.font.size = Pt(size)
    obj.font.bold = bold
    obj.font.italic = False
    obj.font.color.rgb = RGBColor.from_string(color)
    rpr = obj._element.get_or_add_rPr()
    for k in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rpr.rFonts.set(qn('w:' + k), name)


def remove_props(element, tags):
    for tag in tags:
        for child in list(element.findall(qn('w:' + tag))):
            element.remove(child)


for s in doc.sections:
    s.left_margin = s.right_margin = Inches(.8)
    s.top_margin = Inches(.7)
    s.bottom_margin = Inches(.65)
    s.header_distance = s.footer_distance = Inches(.3)
    s.different_first_page_header_footer = True

for name, size, bold in [('Normal', 11, False), ('Title', 25, True),
                         ('Subtitle', 12, False), ('Heading 1', 16, True),
                         ('Heading 2', 11, True), ('Caption', 9, False)]:
    st = doc.styles[name]
    font(st, size=size, bold=bold)
    pf = st.paragraph_format
    pf.space_before, pf.space_after = Pt(0), Pt(8)
    pf.line_spacing = 1.18
    if st.element.pPr is not None:
        remove_props(st.element.pPr, ['pBdr', 'shd', 'contextualSpacing'])
    if name.startswith('Heading'):
        pf.keep_with_next = True

font(doc.styles['SQL Code'], 'DejaVu Sans Mono', 10)
code_style = doc.styles['SQL Code'].paragraph_format
code_style.space_before = code_style.space_after = Pt(0)
code_style.line_spacing = Pt(13)
code_style.left_indent = Inches(.08)

# Replace empty page-break paragraphs with explicit section starts.
for p in list(doc.paragraphs):
    if not p.text and not p._p.xpath('.//w:drawing'):
        p._p.getparent().remove(p._p)

exercise = 0
sql_keywords = re.compile(r"('[^']*'|\b(?:WITH|AS|SELECT|SUM|AVG|COUNT|COALESCE|ROUND|FROM|INNER|LEFT|JOIN|ON|GROUP|BY|WHERE|ORDER|DESC|ROW_NUMBER|OVER|UNION|ALL|IS|NULL|OPTION|MAXRECURSION|AND)\b)")
for p in doc.paragraphs:
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.18
    pf.keep_with_next = False
    pf.keep_together = True
    pf.page_break_before = False
    remove_props(p._p.get_or_add_pPr(), ['pBdr', 'shd', 'contextualSpacing'])
    for run in p.runs:
        font(run)
    if p.style.name == 'Title':
        pf.space_before, pf.space_after = Pt(95), Pt(16)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs: font(run, size=25, bold=True)
    elif p.text.startswith('Northwind'):
        p.style = doc.styles['Subtitle']
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs: font(run, size=12)
    elif p.style.name == 'Heading 1':
        exercise += 1
        label, title = p.text.split('  ', 1)
        p.clear()
        font(p.add_run(label), size=10, bold=True)
        p.add_run().add_break()
        font(p.add_run(title), size=16, bold=True)
        pf.page_break_before = True
        pf.space_after = Pt(12)
        pf.keep_with_next = True
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(100 + exercise))
        start.set(qn('w:name'), f'exercise_{exercise}')
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'), str(100 + exercise))
        p._p.insert(1, start)
        p._p.append(end)
    elif p.style.name == 'Heading 2':
        pf.space_before, pf.space_after = Pt(12), Pt(7)
        pf.keep_with_next = True
        if p.text == 'Асуулгын үр дүн':
            p.text = f'Дасгал {exercise}  Асуулгын үр дүн'
            pf.page_break_before = True
            pf.space_before, pf.space_after = Pt(0), Pt(10)
            for run in p.runs: font(run, size=13, bold=True)
        else:
            for run in p.runs: font(run, size=11, bold=True)
    elif p.style.name == 'SQL Code':
        text = p.text
        p.clear()
        for token in sql_keywords.split(text):
            if not token: continue
            keyword = bool(sql_keywords.fullmatch(token)) and not token.startswith("'")
            font(p.add_run(token), 'DejaVu Sans Mono', 10, keyword,
                 '214F73' if keyword else '000000')
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = Pt(13)
        pf.left_indent = Inches(.08)
        pf.keep_with_next = True
    elif p.style.name == 'Caption':
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.space_before, pf.space_after = Pt(5), Pt(0)
        pf.line_spacing = 1.1
        for run in p.runs: font(run, size=9, color='444444')
    elif p._p.xpath('.//w:drawing'):
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = 1
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for item in p._p.xpath('.//wp:docPr'):
            item.set('descr', f'Дасгал {exercise} асуулгын үр дүнгийн дэлгэцийн зураг')

# Let the final SQL line end the keep-with-next chain.
paragraphs = doc.paragraphs
for i, p in enumerate(paragraphs):
    if i > 0 and paragraphs[i-1].style.name == 'Heading 1':
        p.paragraph_format.space_before = Pt(12)
    if p.style.name == 'SQL Code' and (i+1 == len(paragraphs) or paragraphs[i+1].style.name != 'SQL Code'):
        p.paragraph_format.keep_with_next = False

widths = [[1.55, 2.65, 2.7], [1.15, 1.85, 1.85, 2.05],
          [1.0, 3.55, 1.15, 1.2], [1.05, 1.1, 1.2, 2.9, .65],
          [1.15, 3.35, 1.15, 1.25], [1.55, 2.65, 2.7]]
for table, ws in zip(doc.tables, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    pr = table._tbl.tblPr
    remove_props(pr, ['tblBorders', 'tblCellMar', 'tblInd'])
    pr.find(qn('w:tblW')).set(qn('w:w'), str(round(6.9 * 1440)))
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + edge)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            el.set(qn('w:' + key), value)
        borders.append(el)
    pr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge, value in [('top', '50'), ('bottom', '50'), ('left', '100'), ('right', '100')]:
        el = OxmlElement('w:' + edge)
        el.set(qn('w:w'), value)
        el.set(qn('w:type'), 'dxa')
        margins.append(el)
    pr.append(margins)
    for column, width in zip(table.columns, ws): column.width = Inches(width)
    for ri, row in enumerate(table.rows):
        row.height = None
        row.height_rule = None
        trpr = row._tr.get_or_add_trPr()
        remove_props(trpr, ['trHeight', 'cantSplit', 'tblHeader'])
        trpr.append(OxmlElement('w:cantSplit'))
        if ri == 0: trpr.append(OxmlElement('w:tblHeader'))
        for ci, (cell, width) in enumerate(zip(row.cells, ws)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            remove_props(tcpr, ['shd', 'tcMar', 'tcBorders'])
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '214F73' if ri == 0 else ('F1F5F8' if ri % 2 == 0 else 'FFFFFF'))
            tcpr.append(shade)
            for p in cell.paragraphs:
                pf = p.paragraph_format
                pf.space_before = pf.space_after = Pt(0)
                pf.line_spacing = Pt(12)
                pf.keep_with_next = ri < len(table.rows)-1
                pf.keep_together = True
                value = p.text
                numeric = bool(re.fullmatch(r'[\d,.]+', value))
                p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if ci == 0 and value.isdigit()
                               else WD_ALIGN_PARAGRAPH.RIGHT if numeric
                               else WD_ALIGN_PARAGRAPH.LEFT)
                if ri == 0: p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs: font(run, size=9.5 if len(ws) == 5 else 10, bold=ri == 0,
                                         color='FFFFFF' if ri == 0 else '000000')

for shape in doc.inline_shapes:
    ratio = shape.height / shape.width
    shape.width = Inches(6.9)
    shape.height = int(shape.width * ratio)

# Use the otherwise empty cover for a compact contents list.
first_heading = next(p for p in doc.paragraphs if p.style.name == 'Heading 1')
contents = first_heading.insert_paragraph_before('Агуулга')
contents.paragraph_format.space_before = Pt(65)
contents.paragraph_format.space_after = Pt(14)
for r in contents.runs: font(r, size=12, bold=True)
titles = [p.text.replace('\n', '  ') for p in doc.paragraphs if p.style.name == 'Heading 1']
for i, title in enumerate(titles):
    p = first_heading.insert_paragraph_before(f'{title}\t')
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.9), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.2
    for r in p.runs: font(r, size=10)
    ref = OxmlElement('w:fldSimple')
    ref.set(qn('w:instr'), f' PAGEREF exercise_{i+1} \\h ')
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = str(2 + 2*i)
    r.append(t)
    ref.append(r)
    p._p.append(ref)

for section in doc.sections:
    for p in section.first_page_footer.paragraphs:
        p.clear()
    footer = section.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for value, field in [('CTE даалгавар 1     ', False), ('PAGE', True), (' / ', False), ('NUMPAGES', True)]:
        run = footer.add_run('' if field else value)
        font(run, size=9, color='444444')
        if field:
            el = OxmlElement('w:fldSimple')
            el.set(qn('w:instr'), value)
            run._r.addnext(el)

doc.core_properties.title = 'CTE даалгавар 1'
doc.core_properties.subject = 'Northwind өгөгдлийн сангийн лабораторийн ажил'
assert [p.text for p in doc.paragraphs if p.style.name == 'SQL Code'] == original_sql
assert [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables] == original_tables
assert len(doc.inline_shapes) == 6
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
